import os
import logging
import uuid

logger = logging.getLogger('ai')


class DocumentProcessor:
    """Processes uploaded documents and extracts structured data."""

    MAX_FILE_SIZE = int(os.environ.get('AI_MAX_UPLOAD_SIZE_MB', '10')) * 1024 * 1024
    ALLOWED_EXTENSIONS = {'pdf', 'docx', 'txt'}

    def process(self, file_path: str, file_type: str) -> dict:
        """Process a document and return extracted content."""
        if file_type not in self.ALLOWED_EXTENSIONS:
            raise ValueError(f"Unsupported file type: {file_type}")

        processors = {
            'pdf': self._process_pdf,
            'docx': self._process_docx,
            'txt': self._process_txt,
        }
        return processors[file_type](file_path)

    def _process_pdf(self, file_path: str) -> dict:
        try:
            import PyPDF2

            text_parts = []
            with open(file_path, 'rb') as f:
                reader = PyPDF2.PdfReader(f)
                for page in reader.pages:
                    text = page.extract_text()
                    if text:
                        text_parts.append(text)
            content = '\n\n'.join(text_parts)
            return {
                'content': content,
                'pages': len(text_parts),
                'extraction_method': 'PyPDF2',
            }
        except ImportError:
            logger.error("PyPDF2 not installed. Cannot process PDF files.")
            raise RuntimeError("PDF processing not available")

    def _process_docx(self, file_path: str) -> dict:
        try:
            from docx import Document

            doc = Document(file_path)
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            tables_data = []
            for table in doc.tables:
                for row in table.rows:
                    row_data = [cell.text for cell in row.cells]
                    tables_data.append(row_data)
            content = '\n\n'.join(paragraphs)
            if tables_data:
                content += '\n\nTables:\n'
                for row in tables_data:
                    content += ' | '.join(row) + '\n'
            return {
                'content': content,
                'paragraphs': len(paragraphs),
                'tables': len(doc.tables),
                'extraction_method': 'python-docx',
            }
        except ImportError:
            logger.error("python-docx not installed. Cannot process DOCX files.")
            raise RuntimeError("DOCX processing not available")

    def _process_txt(self, file_path: str) -> dict:
        with open(file_path, 'r', encoding='utf-8', errors='replace') as f:
            content = f.read()
        return {
            'content': content,
            'length': len(content),
            'extraction_method': 'plaintext',
        }

    def sanitize_content(self, content: str) -> str:
        """Remove potentially dangerous content from extracted text."""
        import re

        content = re.sub(r'<script[^>]*>.*?</script>', '', content, flags=re.DOTALL | re.IGNORECASE)
        content = re.sub(r'<[^>]+>', '', content)
        content = re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]', '', content)
        return content.strip()
