import re
import os
import logging
from pathlib import Path

logger = logging.getLogger('ai')

MAX_FILE_SIZE = int(os.environ.get('AI_MAX_UPLOAD_SIZE_MB', '10')) * 1024 * 1024
ALLOWED_EXTENSIONS = {'pdf', 'docx', 'txt'}


def validate_file(filename: str, file_size: int) -> None:
    ext = Path(filename).suffix.lower().lstrip('.')
    if ext not in ALLOWED_EXTENSIONS:
        raise ValueError(f"Unsupported file type: {ext}. Allowed: {', '.join(sorted(ALLOWED_EXTENSIONS))}")
    if file_size > MAX_FILE_SIZE:
        max_mb = MAX_FILE_SIZE // (1024 * 1024)
        raise ValueError(f"File too large: {file_size} bytes. Maximum is {max_mb}MB.")


def extract_pdf(file_path: str) -> str:
    try:
        import PyPDF2
        parts = []
        with open(file_path, 'rb') as f:
            reader = PyPDF2.PdfReader(f)
            for page in reader.pages:
                text = page.extract_text()
                if text:
                    parts.append(text)
        return '\n\n'.join(parts)
    except ImportError:
        raise RuntimeError("PyPDF2 not installed")


def extract_docx(file_path: str) -> str:
    try:
        from docx import Document
        doc = Document(file_path)
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        tables = []
        for table in doc.tables:
            rows = [' | '.join(cell.text for cell in row.cells) for row in table.rows]
            tables.append('\n'.join(rows))
        text = '\n\n'.join(paragraphs)
        if tables:
            text += '\n\nTables:\n' + '\n\n'.join(tables)
        return text
    except ImportError:
        raise RuntimeError("python-docx not installed")


def extract_txt(file_path: str) -> str:
    with open(file_path, 'r', encoding='utf-8', errors='replace') as f:
        return f.read()


EXTRACTORS = {
    'pdf': extract_pdf,
    'docx': extract_docx,
    'txt': extract_txt,
}


def extract(file_path: str, file_type: str) -> str:
    extractor = EXTRACTORS.get(file_type)
    if not extractor:
        raise ValueError(f"No extractor for file type: {file_type}")
    return extractor(file_path)
